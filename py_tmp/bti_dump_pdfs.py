#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
BTI PDF dumper: CSV (колонка «путь») -> один дамп на PDF.

По умолчанию OCR ВЫКЛЮЧЕН: сначала соберите быстрый нативный дамп всего списка.
Затем используйте run_summary.csv и manifest.json, чтобы решить, какие PDF
действительно отправлять на OCR отдельным запуском (--ocr auto / --ocr force).

На каждый PDF создаёт:
  <outdir>/<safe_pdf_name>__<hash>/dump.xlsx
  <outdir>/<safe_pdf_name>__<hash>/text.txt
  <outdir>/<safe_pdf_name>__<hash>/manifest.json

Зависимости:
  py -m pip install pdfplumber pymupdf pytesseract pillow openpyxl tqdm

Tesseract ожидается рядом со скриптом в папке tesseract, например:
  .\tesseract\tesseract.exe
  .\tesseract\tessdata\rus.traineddata

Примеры:
  # Быстрый первый проход: без OCR
  py bti_dump_pdfs.py --csv bti_candidates.csv --outdir bti_dumps

  # Позднее: обработать только PDF, которым нужен OCR
  py bti_dump_pdfs.py --csv ocr_candidates.csv --outdir bti_dumps --ocr auto --overwrite
"""

import argparse
import csv
import hashlib
import json
import os
import re
import unicodedata
from datetime import datetime
from pathlib import Path

import fitz
import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from PIL import Image
import pytesseract

try:
    from tqdm import tqdm
except ImportError:
    def tqdm(items, **kwargs):
        return items


KEYWORDS_RE = re.compile(
    r"эксплик|площад|помещен|квартир|нежил|общая\s+площад|жилая\s+площад",
    re.IGNORECASE,
)
PLAN_WORDS_RE = re.compile(
    r"поэтажн|план\s+(?:этажа|помещен)|схема|масштаб|условн(?:ые|ых)\s+обознач",
    re.IGNORECASE,
)
CYRILLIC_RE = re.compile(r"[А-Яа-яЁё]")
NUMBER_RE = re.compile(r"\b\d{1,4}(?:[,.]\d{1,3})?\b")
SPACE_SPLIT_RE = re.compile(r"\s{2,}")


def norm_text(text):
    text = (text or "").replace("\xa0", " ").replace("\u202f", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def native_is_bad(text):
    text = norm_text(text)
    if len(text) < 80:
        return True
    cid = text.casefold().count("(cid:")
    cyr = len(CYRILLIC_RE.findall(text))
    return cid >= 5 or cyr < max(8, len(text) // 80)


def safe_name(text):
    text = unicodedata.normalize("NFKC", text)
    text = re.sub(r'[<>:"/\\|?*]', "_", text)
    text = re.sub(r"\s+", " ", text).strip(" ._")
    return text[:80] or "pdf"


def short_hash(path):
    return hashlib.sha1(str(path).encode("utf-8", errors="replace")).hexdigest()[:10]


def read_paths(csv_path, column, delimiter):
    result = []
    with csv_path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f, delimiter=delimiter)
        if not reader.fieldnames or column not in reader.fieldnames:
            raise SystemExit(f"В CSV нет колонки «{column}». Найдены: {reader.fieldnames}")
        seen = set()
        for row in reader:
            value = (row.get(column) or "").strip().strip('"')
            if value and value not in seen:
                seen.add(value)
                result.append(Path(value))
    return result


def find_tesseract(script_dir, supplied):
    if supplied:
        p = Path(supplied)
    else:
        candidates = [
            script_dir / "tesseract" / "tesseract.exe",
            script_dir / "Tesseract-OCR" / "tesseract.exe",
            script_dir / "tesseract.exe",
        ]
        p = next((x for x in candidates if x.is_file()), None)
    return p if p and p.is_file() else None


def setup_tesseract(exe, lang):
    if exe is None:
        return False, "tesseract.exe не найден"
    pytesseract.pytesseract.tesseract_cmd = str(exe)
    tessdata = exe.parent / "tessdata"
    if tessdata.is_dir():
        os.environ.setdefault("TESSDATA_PREFIX", str(tessdata))
    try:
        langs = set(pytesseract.get_languages(config=""))
    except Exception as exc:
        return False, f"не удалось запустить Tesseract: {exc}"
    missing = [x for x in lang.split("+") if x not in langs]
    return (not missing), ("" if not missing else f"в Tesseract нет языков: {', '.join(missing)}")


def page_meta(page):
    text = norm_text(page.get_text("text", sort=True))
    blocks = page.get_text("blocks", sort=True)
    image_blocks = [b for b in blocks if len(b) >= 7 and b[6] == 1]
    text_blocks = [b for b in blocks if len(b) >= 7 and b[6] == 0 and str(b[4]).strip()]
    image_area = sum(max(0, b[2] - b[0]) * max(0, b[3] - b[1]) for b in image_blocks)
    page_area = max(1, page.rect.width * page.rect.height)
    try:
        drawings = len(page.get_drawings())
    except Exception:
        drawings = 0
    return {
        "native_text": text,
        "native_chars": len(text),
        "native_bad": native_is_bad(text),
        "text_blocks": len(text_blocks),
        "image_blocks": len(image_blocks),
        "image_share": round(min(1.0, image_area / page_area), 3),
        "image_refs": len(page.get_images(full=False)),
        "drawings": drawings,
        "width": round(page.rect.width, 1),
        "height": round(page.rect.height, 1),
    }


def choose_mode(metas, probe_pages):
    sample = metas[: min(probe_pages, len(metas))]
    if not sample:
        return "unknown"
    good = sum(not x["native_bad"] for x in sample)
    if good == len(sample):
        return "native"
    if good == 0:
        return "scan"
    return "mixed"


def plan_score(meta):
    score, signals = 0, []
    text = meta["native_text"]
    if meta["native_chars"] < 35:
        score += 25
        signals.append("мало текста")
    if meta["image_share"] >= 0.55:
        score += 30
        signals.append("крупное изображение")
    if meta["drawings"] >= 80:
        score += 30
        signals.append("много векторной графики")
    elif meta["drawings"] >= 25:
        score += 10
        signals.append("векторная графика")
    if PLAN_WORDS_RE.search(text):
        score += 20
        signals.append("слова плана")
    if KEYWORDS_RE.search(text):
        score -= 35
        signals.append("слова таблицы площадей")
    return max(0, score), "; ".join(signals)


def render_page(page, dpi):
    scale = dpi / 72
    pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), colorspace=fitz.csRGB, alpha=False)
    return Image.frombytes("RGB", [pix.width, pix.height], pix.samples)


def ocr_page(page, dpi, lang, psm):
    image = render_page(page, dpi)
    try:
        return norm_text(pytesseract.image_to_string(image, lang=lang, config=f"--oem 1 --psm {psm}"))
    finally:
        image.close()


def loose_table(text):
    rows = []
    for line in (text or "").splitlines():
        line = line.strip()
        if not line:
            continue
        cells = [x.strip() for x in SPACE_SPLIT_RE.split(line) if x.strip()]
        if len(cells) >= 2:
            rows.append(cells)
    return rows


def clean_cell(value):
    return norm_text(str(value)) if value is not None else ""


def table_score(rows, context):
    flat = " ".join(clean_cell(c) for row in rows for c in row)
    score, signals = 0, []
    for label, pts, pat in [
        ("экспликация", 50, r"эксплик"),
        ("площадь", 25, r"площад"),
        ("помещение", 15, r"помещен"),
        ("квартира/нежилое", 10, r"квартир|нежил"),
    ]:
        if re.search(pat, flat + " " + context, re.I):
            score += pts
            signals.append(label)
    if len(NUMBER_RE.findall(flat)) >= 8:
        score += 10
        signals.append("много чисел")
    return score, "; ".join(signals)


def context_for_page(text):
    lines = [x.strip() for x in (text or "").splitlines() if x.strip()]
    return " | ".join(lines[:12])[:1200]


def extract_native_tables(pdf_path, page_no):
    try:
        with pdfplumber.open(pdf_path) as pdf:
            tables = pdf.pages[page_no - 1].extract_tables()
    except Exception:
        return []
    return [[list(map(clean_cell, row)) for row in table] for table in tables if table]


def add_table(tables, page_no, method, rows, context):
    rows = [list(row) for row in rows if any(clean_cell(c) for c in row)]
    if not rows:
        return
    ncols = max((len(row) for row in rows), default=0)
    rows = [row + [""] * (ncols - len(row)) for row in rows]
    score, signals = table_score(rows, context)
    tables.append({
        "table_id": f"T{len(tables)+1:05d}", "page": page_no, "method": method,
        "rows": rows, "nrows": len(rows), "ncols": ncols,
        "nonempty": sum(bool(clean_cell(c)) for row in rows for c in row),
        "context": context, "area_score": score, "area_signals": signals,
    })


def write_xlsx(path, pages, tables):
    wb = Workbook()
    ws = wb.active
    ws.title = "index"
    headers = ["table_id", "page", "method", "rows", "cols", "nonempty_cells", "area_score", "area_signals", "context_before"]
    ws.append(headers)
    for t in tables:
        ws.append([t["table_id"], t["page"], t["method"], t["nrows"], t["ncols"], t["nonempty"], t["area_score"], t["area_signals"], t["context"]])

    long_ws = wb.create_sheet("tables_long")
    long_ws.append(["table_id", "page", "method", "row", "col", "value"])
    for t in tables:
        for r_idx, row in enumerate(t["rows"], 1):
            for c_idx, value in enumerate(row, 1):
                long_ws.append([t["table_id"], t["page"], t["method"], r_idx, c_idx, clean_cell(value)])

    wide_ws = wb.create_sheet("tables_wide")
    for t in tables:
        wide_ws.append([f"=== {t['table_id']} | page {t['page']} | {t['method']} | score {t['area_score']} ==="])
        wide_ws.append([t["context"]])
        for row in t["rows"]:
            wide_ws.append([clean_cell(x) for x in row])
        wide_ws.append([])

    text_ws = wb.create_sheet("text_by_page")
    text_ws.append(["page", "text_method", "text", "ocr_action", "skip_reason"])
    for p in pages:
        text_ws.append([p["page"], p["text_method"], p["final_text"], p["ocr_action"], p["skip_reason"]])

    meta_ws = wb.create_sheet("page_index")
    meta_ws.append(["page", "native_chars", "native_bad", "text_blocks", "image_blocks", "image_share", "image_refs", "drawings", "width_pt", "height_pt", "plan_suspect_score", "plan_signals", "ocr_action", "ocr_text_chars", "skip_reason"])
    for p in pages:
        meta_ws.append([p["page"], p["native_chars"], p["native_bad"], p["text_blocks"], p["image_blocks"], p["image_share"], p["image_refs"], p["drawings"], p["width"], p["height"], p["plan_score"], p["plan_signals"], p["ocr_action"], len(p["ocr_text"]), p["skip_reason"]])

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    text_ws.column_dimensions["C"].width = 100
    text_ws.column_dimensions["E"].width = 55
    ws.column_dimensions["I"].width = 100
    wb.save(path)


def write_text(path, pages):
    with path.open("w", encoding="utf-8") as f:
        for p in pages:
            f.write(f"===== PAGE {p['page']:04d} | {p['text_method']} =====\n")
            if p["final_text"]:
                f.write(p["final_text"] + "\n")
            elif p["skip_reason"]:
                f.write(f"[SKIPPED: {p['skip_reason']}]\n")
            f.write("\n")


def process_pdf(pdf_path, outdir, args, ocr_ready):
    folder = outdir / f"{safe_name(pdf_path.stem)}__{short_hash(pdf_path)}"
    manifest_path = folder / "manifest.json"
    if manifest_path.is_file() and not args.overwrite:
        return "skipped_existing", str(folder), {}
    folder.mkdir(parents=True, exist_ok=True)
    manifest = {"source_pdf": str(pdf_path), "created_at": datetime.now().isoformat(timespec="seconds"), "status": "error", "ocr_mode": args.ocr, "ocr_ready": ocr_ready}
    if not pdf_path.is_file():
        manifest["error"] = "source PDF not found"
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        return "missing", str(folder), manifest
    try:
        doc = fitz.open(pdf_path)
    except Exception as exc:
        manifest["error"] = f"cannot open PDF: {type(exc).__name__}: {exc}"
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        return "error", str(folder), manifest
    try:
        metas = [page_meta(doc.load_page(i)) for i in range(doc.page_count)]
        mode = choose_mode(metas, args.probe_pages)
        pages, tables = [], []
        for index, meta in enumerate(metas):
            page_no = index + 1
            page = doc.load_page(index)
            pscore, psignals = plan_score(meta)
            meta.update({"page": page_no, "plan_score": pscore, "plan_signals": psignals, "ocr_action": "not_requested", "skip_reason": "", "ocr_text": "", "final_text": meta["native_text"], "text_method": "native" if meta["native_text"] else "none"})
            needs_ocr = args.ocr == "force" or (args.ocr == "auto" and (mode == "scan" or (mode == "mixed" and meta["native_bad"])))
            has_keywords = bool(KEYWORDS_RE.search(meta["native_text"]))
            if needs_ocr:
                if pscore >= args.plan_skip_score and not has_keywords:
                    meta.update({"ocr_action": "skipped_plan_suspect", "skip_reason": f"plan score {pscore}: {psignals}", "final_text": "", "text_method": "skipped"})
                elif not ocr_ready:
                    meta.update({"ocr_action": "skipped_no_tesseract", "skip_reason": "Tesseract unavailable"})
                else:
                    try:
                        preview = ocr_page(page, args.preview_dpi, args.ocr_lang, args.ocr_psm)
                        if len(preview) < args.preview_min_chars and not KEYWORDS_RE.search(preview):
                            meta.update({"ocr_action": "skipped_low_text_preview", "skip_reason": f"OCR preview: {len(preview)} chars", "final_text": preview, "text_method": "ocr_preview"})
                        else:
                            full = ocr_page(page, args.ocr_dpi, args.ocr_lang, args.ocr_psm)
                            meta.update({"ocr_action": "full_ocr", "ocr_text": full, "final_text": full, "text_method": "ocr"})
                    except Exception as exc:
                        meta.update({"ocr_action": "ocr_error", "skip_reason": f"{type(exc).__name__}: {exc}"})
            context = context_for_page(meta["final_text"])
            native_tables = extract_native_tables(pdf_path, page_no) if not meta["native_bad"] else []
            if native_tables:
                for rows in native_tables:
                    add_table(tables, page_no, "pdfplumber", rows, context)
            else:
                loose = loose_table(meta["final_text"])
                if loose:
                    add_table(tables, page_no, "ocr_loose" if meta["text_method"].startswith("ocr") else "native_loose", loose, context)
            pages.append(meta)
        write_xlsx(folder / "dump.xlsx", pages, tables)
        write_text(folder / "text.txt", pages)
        manifest.update({
            "status": "ok", "document_mode": mode, "page_count": len(pages), "tables_count": len(tables),
            "native_pages": sum(p["text_method"] == "native" for p in pages),
            "ocr_pages": sum(p["text_method"] == "ocr" for p in pages),
            "ocr_preview_pages": sum(p["text_method"] == "ocr_preview" for p in pages),
            "native_bad_pages": sum(p["native_bad"] for p in pages),
            "skipped_plan_pages": sum(p["ocr_action"] == "skipped_plan_suspect" for p in pages),
            "output_xlsx": str(folder / "dump.xlsx"), "output_text": str(folder / "text.txt"),
        })
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        return "ok", str(folder), manifest
    except Exception as exc:
        manifest["error"] = f"{type(exc).__name__}: {exc}"
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        return "error", str(folder), manifest
    finally:
        doc.close()


def main():
    ap = argparse.ArgumentParser(description="Полный дамп текста и таблиц БТИ-PDF из CSV путей.")
    ap.add_argument("--csv", required=True, help="CSV с колонкой «путь»")
    ap.add_argument("--path-column", default="путь")
    ap.add_argument("--sep", default=",", help="Разделитель CSV: ',' или ';'")
    ap.add_argument("--outdir", default="bti_dumps")
    ap.add_argument("--ocr", choices=["off", "auto", "force"], default="off", help="off (по умолчанию), auto или force")
    ap.add_argument("--tesseract-cmd", default=None)
    ap.add_argument("--ocr-lang", default="rus+eng")
    ap.add_argument("--ocr-dpi", type=int, default=300)
    ap.add_argument("--preview-dpi", type=int, default=120)
    ap.add_argument("--ocr-psm", type=int, default=6)
    ap.add_argument("--probe-pages", type=int, default=4)
    ap.add_argument("--preview-min-chars", type=int, default=40)
    ap.add_argument("--plan-skip-score", type=int, default=55)
    ap.add_argument("--overwrite", action="store_true")
    args = ap.parse_args()
    if len(args.sep) != 1:
        raise SystemExit("--sep должен быть одним символом")
    csv_path = Path(args.csv)
    if not csv_path.is_file():
        raise SystemExit(f"Нет CSV: {csv_path}")
    paths = read_paths(csv_path, args.path_column, args.sep)
    if not paths:
        raise SystemExit("В CSV нет непустых путей")
    script_dir = Path(__file__).resolve().parent
    tesseract = find_tesseract(script_dir, args.tesseract_cmd)
    ocr_ready, note = setup_tesseract(tesseract, args.ocr_lang) if args.ocr != "off" else (False, "OCR выключен параметром --ocr off")
    print(f"OCR: {args.ocr}. {note or ('Tesseract: ' + str(tesseract))}")
    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    summary_path = outdir / "run_summary.csv"
    fields = ["source_pdf", "status", "output_folder", "document_mode", "page_count", "tables_count", "native_pages", "native_bad_pages", "ocr_pages", "ocr_preview_pages", "skipped_plan_pages", "ocr_mode"]
    counts = {"ok": 0, "error": 0, "missing": 0, "skipped_existing": 0}
    with summary_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for path in tqdm(paths, desc="BTI PDF", unit="pdf"):
            status, folder, manifest = process_pdf(path, outdir, args, ocr_ready)
            counts[status] = counts.get(status, 0) + 1
            writer.writerow({"source_pdf": str(path), "status": status, "output_folder": folder, **{k: manifest.get(k, "") for k in fields if k not in {"source_pdf", "status", "output_folder"}}})
            f.flush()
    print("Готово:", ", ".join(f"{k}={v}" for k, v in counts.items()))
    print(f"Сводка: {summary_path}")


if __name__ == "__main__":
    main()
