#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Интерпретирует OCR TSV БТИ-дампов и пишет Excel ТОЙ ЖЕ структуры, что native dump:

  index          — одна строка на «таблицу» / визуальный блок страницы
  tables_long    — table_id, page, method, row, col, value
  tables_wide    — блоки подряд в читаемом виде
  text_by_page   — OCR-текст и упорядоченный TSV-текст постранично
  page_index     — OCR-метаданные и статус страницы

Дополнительный лист только для диагностики:
  area_candidates — вероятные строки помещений/площадей; не влияет на основной
                    унифицированный формат и может быть проигнорирован далее.

Вход: CSV с колонкой «путь». Для каждого пути скрипт читает:
  <dump-root>/<safe_stem>__<hash>/ocr_quality/page_*.tsv
  <dump-root>/<safe_stem>__<hash>/ocr_quality_index.csv

Выход: один XLSX на PDF:
  <dump-root>/<safe_stem>__<hash>/dump_ocr.xlsx

Такая структура совместима с native dump.xlsx, созданным bti_dump_pdfs.py:
следующий парсер может читать index/tables_long/tables_wide/text_by_page/page_index
из обоих файлов одинаково. Отличаются только method: native/pdfplumber против
ocr_tsv_lines, и технические поля в page_index.

Зависимости:
  py -m pip install openpyxl

Пример:
  py bti_interpret_dumps.py --csv test_two_pdfs.csv --sep ";" --dump-root bti_dumps

Полностью пересобрать OCR Excel, даже если он уже есть:
  py bti_interpret_dumps.py --csv test_two_pdfs.csv --sep ";" --dump-root bti_dumps --overwrite
"""

import argparse
import csv
import hashlib
import json
import re
import statistics
import unicodedata
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font


AREA_RE = re.compile(r"(?<!\d)(\d{1,4}[,.]\d{1,3}|\d{1,4})(?!\d)")
ROOM_RE = re.compile(
    r"комнат|кухн|коридор|сануз|туалет|ванн|кладов|гардероб|прихож|холл|"
    r"лоджи|балкон|террас|помещен|кабинет|тамбур|веранд|мастерск|подсоб|"
    r"вспомог|жил(?:ая|ое|ых)?|нежил",
    re.IGNORECASE,
)
HEADER_RE = re.compile(r"эксплик|площад|наименован|№\s*помещ|номер\s*помещ", re.IGNORECASE)
TOTAL_RE = re.compile(r"итого|всего|общая\s+площад|жилая\s+площад", re.IGNORECASE)


def safe_name(text):
    text = unicodedata.normalize("NFKC", text)
    text = re.sub(r'[<>:"/\\|?*]', "_", text)
    text = re.sub(r"\s+", " ", text).strip(" ._")
    return text[:80] or "pdf"


def short_hash(path):
    return hashlib.sha1(str(path).encode("utf-8", errors="replace")).hexdigest()[:10]


def dump_folder_for(pdf_path, dump_root):
    return dump_root / f"{safe_name(pdf_path.stem)}__{short_hash(pdf_path)}"


def read_paths(csv_path, column, delimiter):
    paths, seen = [], set()
    with csv_path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f, delimiter=delimiter)
        if not reader.fieldnames or column not in reader.fieldnames:
            raise SystemExit(f"В CSV нет колонки «{column}». Найдены: {reader.fieldnames}")
        for row in reader:
            value = (row.get(column) or "").strip().strip('"')
            if value and value not in seen:
                paths.append(Path(value))
                seen.add(value)
    return paths


def load_index(path):
    rows = []
    if not path.is_file():
        return rows
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def parse_tsv(tsv_path):
    with tsv_path.open("r", encoding="utf-8", errors="replace", newline="") as f:
        reader = csv.DictReader(f, delimiter="\t")
        words = []
        for row in reader:
            text = (row.get("text") or "").strip()
            if not text or row.get("level") != "5":
                continue
            try:
                left = int(row.get("left", "0"))
                top = int(row.get("top", "0"))
                width = int(row.get("width", "0"))
                height = int(row.get("height", "0"))
                conf = float(row.get("conf", "-1"))
            except ValueError:
                continue
            words.append({
                "text": text, "left": left, "top": top, "width": width,
                "height": height, "conf": conf, "right": left + width,
                "bottom": top + height, "cy": top + height / 2,
            })
    return words


def make_lines(words):
    """Собирает PSM 11 words в визуальные строки: top-to-bottom, left-to-right."""
    if not words:
        return []
    heights = [w["height"] for w in words if w["height"] > 0]
    median_height = statistics.median(heights) if heights else 20
    tolerance = max(8.0, median_height * 0.65)
    words = sorted(words, key=lambda w: (w["cy"], w["left"]))
    lines = []
    for word in words:
        matches = [
            (abs(word["cy"] - line["cy"]), line)
            for line in lines
            if abs(word["cy"] - line["cy"]) <= tolerance
        ]
        if matches:
            _, line = min(matches, key=lambda pair: pair[0])
            line["words"].append(word)
            line["cy"] = sum(w["cy"] for w in line["words"]) / len(line["words"])
            line["top"] = min(line["top"], word["top"])
            line["bottom"] = max(line["bottom"], word["bottom"])
            line["left"] = min(line["left"], word["left"])
            line["right"] = max(line["right"], word["right"])
        else:
            lines.append({
                "cy": word["cy"], "top": word["top"], "bottom": word["bottom"],
                "left": word["left"], "right": word["right"], "words": [word],
            })
    lines.sort(key=lambda line: (line["top"], line["left"]))
    for line_no, line in enumerate(lines, start=1):
        line["words"].sort(key=lambda w: w["left"])
        line["line_no"] = line_no
        line["cells"] = [w["text"] for w in line["words"]]
        line["text"] = " ".join(line["cells"])
        line["mean_conf"] = round(sum(w["conf"] for w in line["words"]) / len(line["words"]), 1)
    return lines


def blank_row(ws):
    ws.append([])


def area_candidate(line, header_context):
    text = line["text"]
    numbers = AREA_RE.findall(text)
    if not numbers:
        return None
    raw_area = numbers[-1]
    try:
        area = float(raw_area.replace(",", "."))
    except ValueError:
        return None
    if not (0 < area <= 50000):
        return None
    score, signals = 0, []
    if ROOM_RE.search(text):
        score += 40
        signals.append("тип помещения")
    if header_context:
        score += 35
        signals.append("контекст площади")
    if len(numbers) >= 2:
        score += 10
        signals.append("несколько чисел")
    if 0.5 <= area <= 2000:
        score += 10
        signals.append("правдоподобная площадь")
    if TOTAL_RE.search(text):
        score -= 20
        signals.append("итоговая строка")
    if score < 35:
        return None
    m = re.match(r"^\s*(\d{1,4})\b", text)
    return {
        "room_no_guess": m.group(1) if m else "",
        "area_raw": raw_area,
        "area_guess": area,
        "score": score,
        "signals": "; ".join(signals),
    }


def style_sheet(ws, widths):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def build_workbook(pdf_path, folder, out_path):
    ocr_dir = folder / "ocr_quality"
    index_rows = load_index(folder / "ocr_quality_index.csv")
    if not index_rows:
        return False, "нет ocr_quality_index.csv"

    wb = Workbook()
    index_ws = wb.active
    index_ws.title = "index"
    index_ws.append([
        "table_id", "page", "method", "rows", "cols", "nonempty_cells",
        "area_score", "area_signals", "context_before",
    ])
    long_ws = wb.create_sheet("tables_long")
    long_ws.append(["table_id", "page", "method", "row", "col", "value"])
    wide_ws = wb.create_sheet("tables_wide")
    text_ws = wb.create_sheet("text_by_page")
    text_ws.append(["page", "text_method", "text", "ocr_action", "skip_reason"])
    page_ws = wb.create_sheet("page_index")
    page_ws.append([
        "page", "native_chars", "native_bad", "text_blocks", "image_blocks",
        "image_share", "image_refs", "drawings", "width_pt", "height_pt",
        "plan_suspect_score", "plan_signals", "ocr_action", "ocr_text_chars",
        "skip_reason",
    ])
    candidates_ws = wb.create_sheet("area_candidates")
    candidates_ws.append([
        "table_id", "page", "row", "room_no_guess", "area_raw", "area_guess",
        "score", "signals", "mean_conf", "line_text", "previous_line", "next_line",
    ])

    table_number = 0
    for meta in index_rows:
        page_value = meta.get("page", "")
        try:
            page_no = int(page_value)
        except ValueError:
            page_no = page_value
        status = meta.get("status", "")
        tsv_name = meta.get("tsv_file", "")
        raw_name = meta.get("text_file", "")
        reading_name = meta.get("reading_order_file", "")
        raw_path = ocr_dir / raw_name if raw_name else None
        reading_path = ocr_dir / reading_name if reading_name else None
        tsv_path = ocr_dir / tsv_name if tsv_name else None
        raw_text = raw_path.read_text(encoding="utf-8", errors="replace") if raw_path and raw_path.is_file() else ""
        reading_text = reading_path.read_text(encoding="utf-8", errors="replace") if reading_path and reading_path.is_file() else ""

        if status == "skipped_too_large":
            text_ws.append([page_no, "skipped", "", status, meta.get("reason", "")])
            page_ws.append([page_no, "", "", "", "", "", "", "", "", "", "", "", status, "", meta.get("reason", "")])
            continue
        if not tsv_path or not tsv_path.is_file():
            text_ws.append([page_no, "missing", raw_text, status or "missing_tsv", meta.get("reason", "")])
            page_ws.append([page_no, "", "", "", "", "", "", "", "", "", "", "", status or "missing_tsv", len(raw_text), meta.get("reason", "")])
            continue

        words = parse_tsv(tsv_path)
        lines = make_lines(words)
        final_text = reading_text or "\n".join(line["text"] for line in lines)
        text_ws.append([page_no, "ocr_tsv_reading_order", final_text, status, meta.get("reason", "")])
        page_ws.append([
            page_no, "", "", "", "", "", "", "", "", "", "", "",
            status, len(raw_text), meta.get("reason", ""),
        ])

        if not lines:
            continue

        # Один table_id = один OCR-визуальный блок страницы.
        # Сейчас весь лист — блок: это прямой аналог raw/loose native-table dump.
        table_number += 1
        table_id = f"T{table_number:05d}"
        ncols = max((len(line["cells"]) for line in lines), default=0)
        nonempty = sum(len(line["cells"]) for line in lines)
        context = " | ".join(line["text"] for line in lines[:12])[:1200]
        header_context = 0
        area_score = 0
        area_signals = []
        page_candidates = []

        for idx, line in enumerate(lines):
            if HEADER_RE.search(line["text"]):
                header_context = 8
            candidate = area_candidate(line, header_context)
            if candidate:
                candidate.update({
                    "table_id": table_id,
                    "page": page_no,
                    "row": line["line_no"],
                    "mean_conf": line["mean_conf"],
                    "line_text": line["text"],
                    "previous_line": lines[idx - 1]["text"] if idx > 0 else "",
                    "next_line": lines[idx + 1]["text"] if idx + 1 < len(lines) else "",
                })
                page_candidates.append(candidate)
                area_score = max(area_score, candidate["score"])
                for signal in candidate["signals"].split("; "):
                    if signal and signal not in area_signals:
                        area_signals.append(signal)
            if header_context:
                header_context -= 1

        index_ws.append([
            table_id, page_no, "ocr_tsv_lines", len(lines), ncols, nonempty,
            area_score, "; ".join(area_signals), context,
        ])
        wide_ws.append([f"=== {table_id} | page {page_no} | ocr_tsv_lines | area score {area_score} ==="])
        wide_ws.append([context])
        for line in lines:
            padded = line["cells"] + [""] * (ncols - len(line["cells"]))
            wide_ws.append(padded)
            for col_no, value in enumerate(line["cells"], start=1):
                long_ws.append([table_id, page_no, "ocr_tsv_lines", line["line_no"], col_no, value])
        blank_row(wide_ws)

        for candidate in page_candidates:
            candidates_ws.append([
                candidate["table_id"], candidate["page"], candidate["row"],
                candidate["room_no_guess"], candidate["area_raw"], candidate["area_guess"],
                candidate["score"], candidate["signals"], candidate["mean_conf"],
                candidate["line_text"], candidate["previous_line"], candidate["next_line"],
            ])

    style_sheet(index_ws, {"A": 12, "B": 9, "C": 22, "D": 10, "E": 10, "F": 16, "G": 12, "H": 35, "I": 110})
    style_sheet(long_ws, {"A": 12, "B": 9, "C": 22, "D": 10, "E": 10, "F": 55})
    wide_ws.freeze_panes = "A1"
    wide_ws.column_dimensions["A"].width = 45
    text_ws.column_dimensions["A"].width = 9
    text_ws.column_dimensions["B"].width = 28
    text_ws.column_dimensions["C"].width = 120
    text_ws.column_dimensions["D"].width = 24
    text_ws.column_dimensions["E"].width = 55
    style_sheet(text_ws, {"A": 9, "B": 28, "C": 120, "D": 24, "E": 55})
    style_sheet(page_ws, {"A": 9, "M": 24, "O": 55})
    style_sheet(candidates_ws, {"A": 12, "B": 9, "C": 9, "D": 14, "E": 14, "F": 14, "G": 10, "H": 35, "I": 12, "J": 100, "K": 70, "L": 70})

    wb.save(out_path)
    return True, f"tables={table_number}"


def main():
    ap = argparse.ArgumentParser(description="Создаёт dump_ocr.xlsx в структуре native dump.xlsx из OCR TSV.")
    ap.add_argument("--csv", required=True, help="CSV с колонкой «путь»")
    ap.add_argument("--path-column", default="путь")
    ap.add_argument("--sep", default=",")
    ap.add_argument("--dump-root", required=True)
    ap.add_argument("--overwrite", action="store_true", help="Пересобрать dump_ocr.xlsx, если он уже существует")
    args = ap.parse_args()

    if len(args.sep) != 1:
        raise SystemExit("--sep должен быть одним символом")
    csv_path = Path(args.csv)
    dump_root = Path(args.dump_root)
    if not csv_path.is_file():
        raise SystemExit(f"Нет CSV: {csv_path}")
    if not dump_root.is_dir():
        raise SystemExit(f"Нет папки дампов: {dump_root}")
    paths = read_paths(csv_path, args.path_column, args.sep)
    if not paths:
        raise SystemExit("В CSV нет непустых путей")

    summary_path = dump_root / "interpret_run_summary.csv"
    with summary_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["source_pdf", "status", "output_xlsx", "note"])
        writer.writeheader()
        for pdf_path in paths:
            folder = dump_folder_for(pdf_path, dump_root)
            out_path = folder / "dump_ocr.xlsx"
            if out_path.is_file() and not args.overwrite:
                result = {"source_pdf": str(pdf_path), "status": "skipped_existing", "output_xlsx": str(out_path), "note": ""}
            else:
                try:
                    ok, note = build_workbook(pdf_path, folder, out_path)
                    result = {"source_pdf": str(pdf_path), "status": "ok" if ok else "missing", "output_xlsx": str(out_path) if ok else "", "note": note}
                except Exception as exc:
                    result = {"source_pdf": str(pdf_path), "status": "error", "output_xlsx": "", "note": f"{type(exc).__name__}: {exc}"}
            writer.writerow(result)
            f.flush()
            print(f"{result['status']}: {pdf_path.name} {result['note']}")
    print(f"Сводка: {summary_path}")


if __name__ == "__main__":
    main()
