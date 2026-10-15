#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Позиционная реконструкция OCR TSV БТИ, без словарей помещений и OCR.
Python 3.10+. pip install openpyxl pillow
python bti_position_grid.py FOLDER --outdir grid_test --overwrite
Директория одного документа или корень всех дампов; исходники не изменяются.
Авто-сетка является гипотезой, не подтверждённой экспликацией.
"""
import argparse, csv, hashlib, json, math, re, statistics, traceback
from pathlib import Path
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter

HEADERS={
 'blocks':['document','page','block','x0','y0','x1','y1','rows','columns','mode','flags'],
 'grid_long':['document','page','block','row','column','cell_raw','numeric_candidates','token_ids','x0','y0','x1','y1','mean_conf','flags'],
 'tokens':['document','page','token_id','raw','x0','y0','x1','y1','confidence','noise','block','row','column'],
 'review':['document','page','block','code','detail']}
PURE_NOISE=re.compile('^[\\s\'"`‘’“”|¦!¡=+_~—–\\-.,:;\\\\/()\\[\\]{}<>*#]+$')
NUM=re.compile(r'(?<!\d)\d+(?:[,.]\d+)?(?!\d)')

def median(v,default):return statistics.median(v) if v else default

def read_words(path):
    words=[];bad=[];pagew=pageh=0
    # Quotes in Tesseract TSV are literal OCR content. Never enable CSV quoting.
    with path.open(encoding='utf-8-sig',errors='replace',newline='') as f:
        reader=csv.reader(f,delimiter='\t',quoting=csv.QUOTE_NONE)
        header=next(reader,[])
        if header[:12]!=['level','page_num','block_num','par_num','line_num','word_num','left','top','width','height','conf','text']:
            raise ValueError('Неизвестный заголовок TSV')
        for physical,parts in enumerate(reader,2):
            try:
                if len(parts)<12:raise ValueError('меньше 12 полей')
                level=int(parts[0]);x,y,w,h=map(int,parts[6:10]);conf=float(parts[10])
                if level==1:pagew=max(pagew,x+w);pageh=max(pageh,y+h)
                if level!=5:continue
                text='\t'.join(parts[11:]).strip()
                if not text:continue
                words.append(dict(id=f'T{physical}',text=text,x=x,y=y,r=x+w,b=y+h,w=w,h=h,cy=y+h/2,conf=conf,noise=bool(PURE_NOISE.fullmatch(text)),assignment=None))
            except (ValueError,IndexError) as e:bad.append(f'{physical}: {e}')
    pagew=max(pagew,max((w['r'] for w in words),default=1));pageh=max(pageh,max((w['b'] for w in words),default=1))
    return words,pagew,pageh,bad

def make_rows(words,factor):
    useful=[w for w in words if not w['noise'] and w['h']>0]
    typical=median([w['h'] for w in useful if 6<=w['h']<=120],32)
    tolerance=max(3,typical*factor)
    rows=[]
    for w in sorted(useful,key=lambda z:(z['cy'],z['x'])):
        matches=[(abs(w['cy']-r['cy']),r) for r in rows[-12:] if abs(w['cy']-r['cy'])<=min(tolerance,max(3,w['h']*.65))]
        if matches:
            row=min(matches,key=lambda z:z[0])[1];row['words'].append(w);row['cy']=median([z['cy'] for z in row['words']],0)
        else:rows.append({'cy':w['cy'],'words':[w]})
    rows.sort(key=lambda r:r['cy'])
    for row in rows:
        row['words'].sort(key=lambda w:w['x'])
        row['top']=min(w['y'] for w in row['words']);row['bottom']=max(w['b'] for w in row['words'])
    return rows,typical

def group_blocks(rows,typical,factor):
    if not rows:return []
    gaps=[b['cy']-a['cy'] for a,b in zip(rows,rows[1:]) if b['cy']>a['cy']]
    step=median([g for g in gaps if g>=typical*.7],typical*1.7)
    threshold=max(typical*2.5,step*factor)
    blocks=[];current=[]
    for row in rows:
        if current and row['cy']-current[-1]['cy']>threshold:
            blocks.append(current);current=[]
        current.append(row)
    if current:blocks.append(current)
    return blocks

def infer_cuts(rows,typical,min_support,max_cols):
    # Candidate gutters are horizontal gaps between boxes, not order of OCR words.
    # Each gutter must be supported by several distinct visual rows.
    gap_min=typical*1.0
    intervals=[]
    for i,row in enumerate(rows):
        for a,b in zip(row['words'],row['words'][1:]):
            if b['x']-a['r']>=gap_min:
                intervals.append((a['r']+typical*.20,b['x']-typical*.20,i))
    if not intervals:return []
    required=max(2,math.ceil(len(rows)*min_support))
    coordinates=sorted({x for a,b,_ in intervals for x in (a,b)})
    scored=[]
    for left,right in zip(coordinates,coordinates[1:]):
        x=(left+right)/2
        support={i for a,b,i in intervals if a<=x<=b}
        crossing={i for i,row in enumerate(rows) if any(w['x']<x<w['r'] for w in row['words'])}
        if len(support)>=required and len(crossing)<=max(1,int(len(rows)*.15)):
            scored.append((left,right,len(support)))
    runs=[]
    for a,b,score in scored:
        if runs and a<=runs[-1][1]+1:runs[-1]=(runs[-1][0],b,max(score,runs[-1][2]))
        else:runs.append((a,b,score))
    runs=[r for r in runs if r[1]-r[0]>=typical*.30]
    runs=sorted(runs,key=lambda r:(r[2],r[1]-r[0]),reverse=True)[:max_cols-1]
    return sorted((a+b)/2 for a,b,_ in runs)

def excel_text(ws,values):
    cells=[]
    for v in values:
        if isinstance(v,str):
            # Do not let OCR tokens starting '=' become Excel formulas.
            v=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',v)[:32700]
            c=WriteOnlyCell(ws,value=v);c.data_type='s';cells.append(c)
        else:cells.append(v)
    ws.append(cells)

class Output:
    def __init__(self,out):
        self.wb=Workbook(write_only=True);self.files=[];self.sheets={};self.counts={};self.parts={};self.writers={}
        for name,h in HEADERS.items():
            f=(out/(name+'.csv')).open('w',encoding='utf-8-sig',newline='');self.files.append(f)
            wr=csv.writer(f);wr.writerow(h);self.writers[name]=wr
            self.newsheet(name,h)
        self.newsheet('grid_wide',['document','page','block','row','y']+[f'C{i:02}' for i in range(1,65)]+['flags'])
    def newsheet(self,name,header):
        part=self.parts.get(name,0)+1;self.parts[name]=part
        ws=self.wb.create_sheet(name if part==1 else f'{name}_{part}');ws.freeze_panes='F2'
        for i in range(1,len(header)+1):ws.column_dimensions[get_column_letter(i)].width=18 if i>5 else 20
        excel_text(ws,header);self.sheets[name]=(ws,header);self.counts[name]=1
    def add(self,name,values):
        if name in self.writers:self.writers[name].writerow(values)
        ws,header=self.sheets[name]
        if self.counts[name]>=1000000:self.newsheet(name,header);ws,_=self.sheets[name]
        excel_text(ws,values);self.counts[name]+=1
    def close(self,out):
        for f in self.files:f.close()
        tmp=out/'grid.tmp.xlsx';self.wb.save(tmp);tmp.replace(out/'grid.xlsx')

def overlay(folder,out,page,words,regions):
    from PIL import Image,ImageDraw
    images=list((folder/'control').glob(f'page{page:04d}_prepared.png'))
    if not images:return
    with Image.open(images[0]) as image:
        image=image.convert('RGB');draw=ImageDraw.Draw(image)
        for box,cuts,rows,bid in regions:
            draw.rectangle(box,outline='blue',width=4)
            for x in cuts:draw.line((x,box[1],x,box[3]),fill='red',width=3)
            for row in rows:draw.line((box[0],row['cy'],box[2],row['cy']),fill='green',width=2)
        for w in words:
            if w['assignment']:draw.rectangle((w['x'],w['y'],w['r'],w['b']),outline='orange',width=1)
        image.thumbnail((2200,3000));image.save(out/f'page{page:04d}_overlay.png')

def process(folder,out,args):
    if (out/'grid.xlsx').exists() and not args.overwrite:return 'cached'
    out.mkdir(parents=True,exist_ok=True);writer=Output(out)
    doc=folder.name
    for path in sorted(folder.glob('page*_result.tsv')):
        m=re.fullmatch(r'page(\d+)_result\.tsv',path.name)
        if not m:continue
        page=int(m[1])
        try:words,pw,ph,bad=read_words(path)
        except Exception as e:writer.add('review',[doc,page,'','tsv_error',str(e)]);continue
        for issue in bad:writer.add('review',[doc,page,'','bad_tsv_line',issue])
        rows,typical=make_rows(words,args.row_tolerance)
        blocks=group_blocks(rows,typical,args.block_gap);regions=[]
        for bi,br in enumerate(blocks,1):
            bid=f'P{page:04d}B{bi:03d}'
            cuts=[v*pw for v in args.fractions] if args.fractions else infer_cuts(br,typical,args.support,args.max_cols)
            mode='manual' if args.fractions else 'auto_gutters'
            flags='unverified_geometry'
            if not cuts:flags+=';no_stable_columns'
            if len(br)<3:flags+=';short_block'
            box=(min(w['x'] for r in br for w in r['words']),min(r['top'] for r in br),max(w['r'] for r in br for w in r['words']),max(r['bottom'] for r in br))
            writer.add('blocks',[doc,page,bid,*box,len(br),len(cuts)+1,mode,flags])
            writer.add('review',[doc,page,bid,'geometry_hypothesis',flags])
            regions.append((box,cuts,br,bid))
            for ri,row in enumerate(br,1):
                slots=defaultdict(list);rflags=[]
                for w in row['words']:
                    cx=(w['x']+w['r'])/2
                    col=sum(cx>=x for x in cuts)+1
                    w['assignment']=(bid,ri,col);slots[col].append(w)
                    if any(w['x']<x<w['r'] for x in cuts):rflags.append('crosses_column_boundary')
                wide=['']*(len(cuts)+1)
                for col,ww in sorted(slots.items()):
                    ww.sort(key=lambda w:w['x']);raw=' '.join(w['text'] for w in ww);wide[col-1]=raw
                    candidates=[]
                    for w in ww:
                        for match in NUM.finditer(w['text']):
                            candidates.append({'raw':match[0],'value':match[0].replace(',','.'),'token':w['id'],'offset':match.start(),'conf':w['conf']})
                    cellflags=[]
                    if min(w['conf'] for w in ww)<args.review_conf:cellflags.append('low_conf')
                    if len(candidates)>1:cellflags.append('multiple_numbers')
                    if any(w['x']<x<w['r'] for w in ww for x in cuts):cellflags.append('crosses_column_boundary')
                    writer.add('grid_long',[doc,page,bid,ri,f'C{col:02}',raw,json.dumps(candidates,ensure_ascii=False),' '.join(w['id'] for w in ww),min(w['x'] for w in ww),min(w['y'] for w in ww),max(w['r'] for w in ww),max(w['b'] for w in ww),round(statistics.mean(w['conf'] for w in ww),2),';'.join(cellflags)])
                writer.add('grid_wide',[doc,page,bid,ri,round(row['cy'],1),*wide,*(['']*(64-len(wide))),';'.join(sorted(set(rflags)))])
        for w in words:
            assignment=w['assignment'] or ('','','')
            writer.add('tokens',[doc,page,w['id'],w['text'],w['x'],w['y'],w['r'],w['b'],w['conf'],w['noise'],*assignment])
        if args.overlays:
            try:overlay(folder,out,page,words,regions)
            except Exception as e:writer.add('review',[doc,page,'','overlay_error',str(e)])
        print(f'  {path.name}: {len(words)} tokens, {len(rows)} rows, {len(blocks)} blocks',flush=True)
    writer.close(out)
    (out/'settings.json').write_text(json.dumps({'source':str(folder.resolve()),'args':{k:str(v) for k,v in vars(args).items()},'note':'Columns are positional hypotheses. No room semantics or verified object boundaries.'},ensure_ascii=False,indent=2),encoding='utf-8')
    return 'ok'

def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('root',type=Path)
    ap.add_argument('--outdir',type=Path,default=Path('bti_position_grid'))
    ap.add_argument('--overwrite',action='store_true')
    ap.add_argument('--row-tolerance',type=float,default=.45,help='Допуск Y относительно медианной высоты символов')
    ap.add_argument('--block-gap',type=float,default=3.0,help='Разрыв блоков относительно типичного шага строк')
    ap.add_argument('--support',type=float,default=.30,help='Доля строк, поддерживающих промежуток колонок')
    ap.add_argument('--max-cols',type=int,default=20)
    ap.add_argument('--column-cuts',default='',help='Ручные границы X как доли ширины: 0.15,0.25,0.45,0.60')
    ap.add_argument('--review-conf',type=float,default=40)
    ap.add_argument('--overlays',action='store_true',help='Наложение сетки на существующие control/*prepared.png; PDF не нужен')
    a=ap.parse_args()
    if not a.root.is_dir():ap.error('Нет входной папки')
    if not 0<a.support<=1 or not 1<=a.max_cols<=64 or a.row_tolerance<=0 or a.block_gap<=0:ap.error('Некорректные параметры')
    try:a.fractions=sorted(set(float(v) for v in a.column_cuts.split(',') if v.strip()))
    except ValueError:ap.error('Неверный --column-cuts')
    if len(a.fractions)>63 or any(not 0<v<1 for v in a.fractions):ap.error('Границы: от 0 до 1, не более 63')
    folders=sorted(set(p.parent for p in a.root.rglob('page*_result.tsv')))
    if not folders:ap.error('Не найдены pageNNNN_result.tsv')
    a.outdir.mkdir(parents=True,exist_ok=True)
    with (a.outdir/'run_summary.csv').open('w',encoding='utf-8-sig',newline='') as f:
        wr=csv.writer(f);wr.writerow(['folder','status','output','error'])
        for folder in folders:
            uid=hashlib.sha1(str(folder.resolve()).encode()).hexdigest()[:8]
            out=a.outdir/(folder.name+'__'+uid)
            try:status=process(folder,out,a);error=''
            except Exception:status='error';error=traceback.format_exc();out.mkdir(parents=True,exist_ok=True);(out/'error.txt').write_text(error,encoding='utf-8')
            wr.writerow([str(folder),status,str(out),error]);f.flush()
            print(status,folder.name,flush=True)
if __name__=='__main__':main()
